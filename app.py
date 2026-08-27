import io
import zipfile
import shutil

import pypdfium2 as pdfium
import pytesseract
import streamlit as st
from pypdf import PdfReader, PdfWriter
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

st.set_page_config(page_title="Conversor PDF OCR", page_icon="📄", layout="centered")

DEFAULT_BLOCK_SIZE = 20


def check_tesseract():
    if not shutil.which("tesseract"):
        st.error("Tesseract no está disponible en el servidor.")
        st.stop()


def get_pdf(pdf_bytes):
    return pdfium.PdfDocument(io.BytesIO(pdf_bytes))


def render_page(pdf, index, dpi):
    scale = dpi / 72
    page = pdf[index]
    bitmap = page.render(scale=scale)
    return bitmap.to_pil().convert("L")


def image_to_ocr_pdf(image, lang):
    return pytesseract.image_to_pdf_or_hocr(image, extension="pdf", lang=lang)


def merge_pdf_pages(pages):
    writer = PdfWriter()
    for pdf_page in pages:
        reader = PdfReader(io.BytesIO(pdf_page))
        for page in reader.pages:
            writer.add_page(page)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def create_word_document():
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    return document


def add_text_to_word(document, text, add_page_break=True):
    for line in text.splitlines():
        line = line.strip()
        if line:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line)
            run.font.name = "Arial"
            run.font.size = Pt(12)
            run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    if add_page_break:
        document.add_page_break()


def word_to_bytes(document):
    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    return output.getvalue()


def process_pdf_ocr(pdf_bytes, dpi, lang, block_size):
    pdf = get_pdf(pdf_bytes)
    total_pages = len(pdf)
    zip_buffer = io.BytesIO()

    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        block_pages = []
        block_number = 1
        progress = st.progress(0)
        status = st.empty()

        for i in range(total_pages):
            status.info(f"Procesando PDF OCR: página {i + 1} de {total_pages}")
            image = render_page(pdf, i, dpi)
            pdf_page = image_to_ocr_pdf(image, lang)
            block_pages.append(pdf_page)

            if len(block_pages) == block_size or i == total_pages - 1:
                merged = merge_pdf_pages(block_pages)
                start_page = i - len(block_pages) + 2
                end_page = i + 1
                filename = f"PDF_OCR_{block_number:03d}_pag_{start_page}_{end_page}.pdf"
                zip_file.writestr(filename, merged)
                block_pages = []
                block_number += 1

            progress.progress((i + 1) / total_pages)

    return zip_buffer.getvalue()


def process_word(pdf_bytes, dpi, lang):
    pdf = get_pdf(pdf_bytes)
    total_pages = len(pdf)
    document = create_word_document()
    progress = st.progress(0)
    status = st.empty()

    for i in range(total_pages):
        status.info(f"Convirtiendo a Word: página {i + 1} de {total_pages}")
        image = render_page(pdf, i, dpi)
        text = pytesseract.image_to_string(image, lang=lang)
        add_text_to_word(document, text, add_page_break=(i < total_pages - 1))
        progress.progress((i + 1) / total_pages)

    return word_to_bytes(document)


def main():
    st.title("📄 Conversor PDF OCR")
    st.write("Convierte documentos PDF escaneados a PDF con OCR, Word editable en Arial 12, o ambos.")

    check_tesseract()

    st.subheader("1. Formato de salida")
    output_format = st.radio(
        "¿Qué deseas obtener?",
        ["PDF con OCR", "Word editable - Arial 12", "PDF OCR + Word"],
        index=0,
    )

    st.subheader("2. Configuración")
    language_options = {
        "Español": "spa",
        "Español + Inglés": "spa+eng",
        "Inglés": "eng",
    }
    language_label = st.selectbox("Idioma del documento", list(language_options.keys()))
    lang = language_options[language_label]

    quality_options = {
        "Baja - más rápido": 100,
        "Media - recomendado": 120,
        "Alta - mejor calidad": 140,
    }
    quality_label = st.selectbox("Calidad", list(quality_options.keys()), index=1)
    dpi = quality_options[quality_label]

    block_size = DEFAULT_BLOCK_SIZE
    if output_format in ["PDF con OCR", "PDF OCR + Word"]:
        block_size = st.slider(
            "Páginas por cada archivo PDF OCR",
            min_value=5,
            max_value=50,
            value=20,
            step=5,
        )
        st.caption("Para documentos muy grandes se recomienda 10 a 20 páginas por archivo.")

    st.subheader("3. Subir documento")
    uploaded_file = st.file_uploader("Selecciona un archivo PDF", type=["pdf"])
    if uploaded_file is None:
        return

    pdf_bytes = uploaded_file.read()

    try:
        pdf = get_pdf(pdf_bytes)
        total_pages = len(pdf)
    except Exception as error:
        st.error(f"No fue posible leer el PDF: {error}")
        return

    size_mb = len(pdf_bytes) / 1024 / 1024
    st.success(f"Archivo cargado · {total_pages} páginas · {size_mb:.2f} MB")

    if total_pages > 500:
        st.warning("El documento es muy grande. El procesamiento puede tardar bastante.")

    st.subheader("4. Convertir")

    if st.button("🚀 Iniciar conversión", use_container_width=True):
        base_name = uploaded_file.name.rsplit(".", 1)[0]

        if output_format == "PDF con OCR":
            result_pdf = process_pdf_ocr(pdf_bytes, dpi, lang, block_size)
            st.success("PDF OCR listo.")
            st.download_button(
                "⬇️ Descargar PDF OCR",
                data=result_pdf,
                file_name=f"{base_name}_PDF_OCR.zip",
                mime="application/zip",
                use_container_width=True,
            )

        elif output_format == "Word editable - Arial 12":
            result_word = process_word(pdf_bytes, dpi, lang)
            st.success("Word editable listo en Arial 12.")
            st.download_button(
                "⬇️ Descargar Word",
                data=result_word,
                file_name=f"{base_name}_OCR.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
            )

        else:
            st.markdown("#### Generando PDF OCR")
            result_pdf = process_pdf_ocr(pdf_bytes, dpi, lang, block_size)

            st.markdown("#### Generando Word")
            result_word = process_word(pdf_bytes, dpi, lang)

            st.success("PDF OCR y Word listos.")
            st.download_button(
                "⬇️ Descargar PDF OCR",
                data=result_pdf,
                file_name=f"{base_name}_PDF_OCR.zip",
                mime="application/zip",
                use_container_width=True,
            )
            st.download_button(
                "⬇️ Descargar Word Arial 12",
                data=result_word,
                file_name=f"{base_name}_OCR.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
            )


if __name__ == "__main__":
    main()
